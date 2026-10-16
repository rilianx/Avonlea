"""Lleva lo editado en planilla/avonlea.xlsx al juego: «Frases» → src/datos/frases.txt, «Escenas» → la cabecera de
cada escena de src/datos/guion.txt, «Noticias» → src/datos/noticias.txt, «Llegadas» → src/datos/llegadas.txt.
Todo o nada: revisa cada fila de cada hoja y, si algo está mal, no cambia ningún archivo y dice la hoja y la fila.
Si todo está bien, escribe los cambios en su lugar (los comentarios, el orden y los diálogos quedan), compila el juego
(el revisor de la historia, tools/grafo.js) y lo abre (tools/verificar_carga.js): si algo falla, vuelve todo atrás.
Uso: python3 tools/importar_planilla.py [planilla/avonlea.xlsx]"""
import json, os, re, shutil, subprocess, sys
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(__file__))
import bloques as B

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, 'planilla/avonlea.xlsx')
FRASES = os.path.join(RAIZ, 'src/datos/frases.txt')
TIPOS = ['dice', 'ana', 'primavera', 'verano', 'otoño', 'invierno']
ESTACIONES = ['primavera', 'verano', 'otoño', 'invierno']
grafo = json.load(open(os.path.join(RAIZ, 'trazas/grafo.json'), encoding='utf-8'))
metas = set(grafo['metas'])

wb = load_workbook(XLSX)   # (the formulas are not needed: only what was typed)
P = wb['Personas']
quienes = set(P.cell(row=r, column=1).value for r in range(2, P.max_row + 1) if P.cell(row=r, column=1).value) | {'ana'}
texto = lambda v: '' if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)).strip()
errores, avisos = [], []


def condiciones_malas(conds):
    """las que no son condiciones del guion (las claves que conoce el revisor de la historia, o una cualidad)"""
    malas = []
    for c in conds:
        k = c.lstrip('!').split(':')[0].split('>')[0].split('<')[0].split('=')[0]
        if not re.match(r'^[\w>=<.:!-]+$', c) or (k not in grafo['claves'] and k not in grafo['cualidades']):
            malas.append(c)
    return malas


lista = lambda v: [x.strip() for x in v.split(',') if x.strip()]

# --- Frases
ws = wb['Frases']
cab = [ws.cell(row=1, column=c).value for c in range(1, 11)]
if cab != ['Quién', 'Nombre', 'Tipo', 'Frase', '¿Conoce a Ana?', 'Después de (hecho o meta)', 'Días después', 'Antes de (hecho o meta)', 'Estación', 'Otras condiciones']:
    sys.exit('La hoja «Frases» no tiene las columnas de siempre (¿se movió o renombró alguna?):\n  ' + ' | '.join(map(str, cab)))
por_quien, orden = {}, []
for r in range(2, ws.max_row + 1):
    q, tipo, frase, conoce, despues, dias, antes, est, otras = (texto(ws.cell(row=r, column=c).value) for c in (1, 3, 4, 5, 6, 7, 8, 9, 10))
    if not any((q, tipo, frase, conoce, despues, dias, antes, est, otras)):
        continue   # (an empty row)
    malo = lambda m: errores.append(f'Frases, fila {r}: {m}')
    if q not in quienes - {'ana'}: malo(f'«{q}» no es nadie del juego (hoja Personas)')
    if tipo not in TIPOS: malo(f'el tipo «{tipo}» no existe (dice, ana, o una estación)')
    if not frase: malo('la frase está vacía')
    if '[' in frase or ']' in frase or '\n' in frase: malo('la frase no puede tener corchetes ni saltos de línea')
    if conoce not in ('', 'sí', 'no'): malo(f'«¿Conoce a Ana?» debe ser sí, no o vacío (dice «{conoce}»)')
    if est and est not in ESTACIONES: malo(f'la estación «{est}» no existe')
    if dias and not despues: malo('hay «Días después» pero no «Después de»')
    if dias and not re.fullmatch(r'\d+', dias): malo(f'«Días después» debe ser un número (dice «{dias}»)')
    for nombre, v in (('Después de', despues), ('Antes de', antes)):
        if v and not re.fullmatch(r'\w+', v): malo(f'«{nombre}»: «{v}» debe ser un solo nombre de hecho o meta, sin espacios')
    conds = []
    if conoce: conds.append(('' if conoce == 'sí' else '!') + 'conoce:' + q)
    if despues:
        if dias and re.fullmatch(r'\d+', dias): conds.append(f'dias_desde:{despues}:{dias}')
        elif not dias: conds.append(('cuaderno:' if despues in metas and despues not in grafo['hechos'] else 'hecho:') + despues)
    if antes: conds.append('!' + ('cuaderno:' if antes in metas and antes not in grafo['hechos'] else 'hecho:') + antes)
    if est: conds.append('estacion:' + est)
    conds += lista(otras)
    for c in condiciones_malas(conds): malo(f'la condición «{c}» no es una condición del guion')
    if q not in por_quien: por_quien[q] = []; orden.append(q)
    por_quien[q].append((tipo, ('[' + ', '.join(conds) + '] ' if conds else '') + frase))
viejo_frases = open(FRASES, encoding='utf-8').read()
cabecera = viejo_frases[:viejo_frases.index('\n===')] if '\n===' in viejo_frases else ''
partes = [cabecera.rstrip('\n'), '']
for q in orden:
    partes += [f'=== {q} ==='] + [f'{tipo}: {linea}' for tipo, linea in por_quien[q]] + ['']
nuevos = {FRASES: '\n'.join(partes)}


# --- Escenas, Noticias, Llegadas: cada fila, los campos de un bloque
def revisar(hoja, id_, pares, malo):
    for k, v in pares:
        if hoja == 'Escenas':
            if k == 'si':
                for c in condiciones_malas(lista(v)): malo(f'«Si…»: «{c}» no es una condición del guion')
            elif k == 'prioridad' and not re.fullmatch(r'-?\d+(\.\d+)?', v): malo(f'la prioridad debe ser un número (dice «{v}»)')
            elif k == 'en' and not all(re.fullmatch(r'\w+=\w+', x) for x in lista(v)): malo(f'«Dónde está cada uno»: «{v}» (debe ser quien=lugar, con comas)')
            elif k in ('interrumpe', 'grande') and v != 'sí': malo(f'«{k}» debe ser «sí» o vacío')
            elif k == 'sujeta':
                for x in lista(v):
                    if x not in quienes: malo(f'«Quedan quietos»: «{x}» no es nadie del juego')
        elif hoja == 'Noticias':
            if k == 'saben':
                for x in lista(v):
                    if x not in quienes: malo(f'«La saben»: «{x}» no es nadie del juego')
            elif k == 'saben si':
                m = re.match(r'^(.+?):\s+(.+)$', v)
                if not m: malo(f'«La saben desde que…»: «{v}» (debe ser «condiciones: quiénes», con un espacio después de los dos puntos)'); continue
                for c in condiciones_malas(lista(m.group(1))): malo(f'«La saben desde que…»: «{c}» no es una condición del guion')
                for x in lista(m.group(2)):
                    if x not in quienes: malo(f'«La saben desde que…»: «{x}» no es nadie del juego')
            elif k == 'llega a':
                m = re.match(r'^(\w+):\s*([\d.]+)$', v)
                if not m or m.group(1) not in quienes: malo(f'«Le llega a»: «{v}» (debe ser «quién: horas»)')
            elif k == 'se cuenta por' and not re.fullmatch(r'\d+', v): malo(f'«Se cuenta por» debe ser un número de días (dice «{v}»)')
            elif k == 'dice de':
                m = re.match(r'^(\w+):\s*(.+)$', v)
                if not m or m.group(1) not in quienes: malo(f'«Cómo la cuenta cada uno»: «{v[:40]}» (debe ser «quién: texto»)')
            elif k == 'al conocer a ana' and v != 'sí': malo('«Al conocer a Ana» debe ser «sí» o vacío')
        elif hoja == 'Llegadas':
            if k in ('llega si', 'se va si'):
                for c in condiciones_malas(lista(v)): malo(f'«{k}»: «{c}» no es una condición del guion')
            elif k == 'de visita' and v != 'sí': malo('«De visita» debe ser «sí» o vacío')
    if hoja == 'Llegadas' and id_ not in quienes: malo(f'«{id_}» no es nadie del juego (hoja Personas)')


for hoja, H in B.HOJAS.items():
    if hoja not in wb.sheetnames:
        continue
    ws = wb[hoja]; ruta = os.path.join(RAIZ, H['archivo'])
    n = len(H['columnas']) + (1 if H['otros'] else 0)
    esperado = [H['id']] + [t for _, t, _ in H['columnas']] + ([H['otros'][0]] if H['otros'] else [])
    cab = [ws.cell(row=1, column=c).value for c in range(1, n + 2)][:n + 1]
    if cab != esperado:
        errores.append(f'{hoja}: las columnas no son las de siempre (¿se movió o renombró alguna?)'); continue
    partes = B.partir(open(ruta, encoding='utf-8').read())
    bloques = {p[1]: i for i, p in enumerate(partes) if p[0] == 'bloque'}
    vistos = set()
    for r in range(2, ws.max_row + 1):
        id_ = texto(ws.cell(row=r, column=1).value)
        valores = [texto(ws.cell(row=r, column=c).value) for c in range(2, n + 2)]
        if not id_ and not any(valores):
            continue
        malo = lambda m, r=r: errores.append(f'{hoja}, fila {r}: {m}')
        if not re.fullmatch(r'\w+', id_ or ''): malo(f'«{id_}» no sirve de nombre (letras, números y _ , sin espacios)'); continue
        if id_ in vistos: malo(f'«{id_}» está dos veces'); continue
        vistos.add(id_)
        pares = B.de_celdas(hoja, valores)
        revisar(hoja, id_, pares, malo)
        if id_ in bloques:
            i = bloques[id_]; p = partes[i]
            partes[i] = (p[0], p[1], B.actualizar(p[2], pares, H['clasificar'], H['escribir']), p[3])
        elif H['nuevas']:
            if not (partes and partes[-1][0] == 'texto' and partes[-1][1] and not partes[-1][1][-1].strip()): partes.append(('texto', ['']))
            partes.append(('bloque', id_, [f'=== {id_} ==='] + [H['escribir'](k, v) for k, v in pares], []))
        else:
            malo(f'no hay una escena «{id_}» en {H["archivo"]} (las escenas nuevas se escriben allí, con sus líneas)')
    for id_ in bloques:
        if id_ not in vistos:
            avisos.append(f'{hoja}: «{id_}» no está en la planilla; queda como estaba en {H["archivo"]} (para borrarlo, bórralo allí)')
    nuevos[ruta] = B.juntar(partes)

if errores:
    sys.exit('No cambié nada. Hay que corregir la planilla:\n  ' + '\n  '.join(errores))
for a in avisos:
    print('  aviso: ' + a)
cambian = {ruta: t for ruta, t in nuevos.items() if t.strip() != open(ruta, encoding='utf-8').read().strip()}
if not cambian:
    print('La planilla y los archivos de src/datos/ ya dicen lo mismo: no hay nada que cambiar.'); sys.exit(0)
malas_antes = set(grafo['malas'])
for ruta, t in cambian.items():
    shutil.copy(ruta, ruta + '.antes')
    open(ruta, 'w', encoding='utf-8').write(t if t.endswith('\n') else t + '\n')


def atras(motivo):
    for ruta in cambian:
        shutil.copy(ruta + '.antes', ruta); os.remove(ruta + '.antes')
    subprocess.run(['node', os.path.join(RAIZ, 'tools/build.js')], capture_output=True)
    sys.exit('No cambié nada: ' + motivo)


subprocess.run(['node', os.path.join(RAIZ, 'tools/build.js')], capture_output=True, text=True)
g = json.load(open(os.path.join(RAIZ, 'trazas/grafo.json'), encoding='utf-8'))
nuevas_malas = [m for m in g['malas'] if m not in malas_antes] + [f'«{h}» se pide y nada lo deja' for h in g['nunca'] if h not in grafo['nunca']]
if nuevas_malas:
    atras('el revisor de la historia encontró problemas:\n  ' + '\n  '.join(nuevas_malas))
v = subprocess.run(['node', os.path.join(RAIZ, 'tools/verificar_carga.js')], capture_output=True, text=True, env={**os.environ, 'NODE_PATH': os.environ.get('NODE_PATH', '/opt/node22/lib/node_modules')})
if v.returncode != 0:
    atras('el juego no carga con estos cambios:\n  ' + (v.stdout + v.stderr).strip().replace('\n', '\n  '))
for ruta in cambian:
    os.remove(ruta + '.antes')
print('→ ' + ', '.join(os.path.relpath(r, RAIZ) for r in cambian) + ' · el juego, compilado y abierto sin errores')
print('  Para ver la planilla al día (y los huecos): node tools/planilla.js')
